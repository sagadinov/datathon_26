"""Dataset labels: open_labels.xlsx parsing, video matching by number N, expected T-type proxy."""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from src.common import ROOT

DATA_DIR = ROOT / "data" / "open"
LABELS_XLSX = DATA_DIR / "open_labels.xlsx"
VIDEOS_DIR = DATA_DIR / "videos"

# Our own type reconstruction (CLAUDE.md section 3); labels have no type column.
TYPE_BY_N: dict[int, str] = {
    **dict.fromkeys([4, 26, 38, 15], "T1"),
    **dict.fromkeys([28, 22, 11], "T2"),
    **dict.fromkeys([20], "T3"),
    **dict.fromkeys([18, 30], "T4"),
    **dict.fromkeys([32, 33, 40], "T5"),
    **dict.fromkeys([46, 47, 24], "T6"),
    **dict.fromkeys([1, 44, 13, 2], "T7"),
    **dict.fromkeys([34], "T8"),
    **dict.fromkeys([17], "T9"),
}


@dataclass
class Label:
    n: int
    name: str
    violation: bool
    t_sec: float | None
    type_id: str | None
    video: Path | None


def video_number(name: str) -> int | None:
    """'X-4', 'Х-4.mov', '#U0425-4.mov' -> 4 (match by number, never by the letter)."""
    m = re.search(r"-(\d+)", Path(name).stem)
    return int(m.group(1)) if m else None


def parse_timestamp(text: str | None) -> float | None:
    """'22 секунда' -> 22.0; '1минута 8 секунда' -> 68.0; '' -> None."""
    if not text:
        return None
    text = str(text).lower()
    minutes = re.search(r"(\d+)\s*мин", text)
    seconds = re.search(r"(\d+)\s*сек", text)
    if not minutes and not seconds:
        nums = re.findall(r"\d+", text)
        return float(nums[0]) if nums else None
    return 60.0 * (int(minutes.group(1)) if minutes else 0) + (int(seconds.group(1)) if seconds else 0)


def index_videos(videos_dir: Path = VIDEOS_DIR) -> dict[int, Path]:
    out: dict[int, Path] = {}
    for p in sorted(videos_dir.glob("*.mov")):
        n = video_number(p.name)
        if n is not None:
            out[n] = p
    return out


def load_labels(xlsx: Path = LABELS_XLSX, videos_dir: Path = VIDEOS_DIR) -> list[Label]:
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["labels"] if "labels" in wb.sheetnames else wb[wb.sheetnames[0]]
    videos = index_videos(videos_dir)
    labels: list[Label] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        name = str(row[0]).strip()
        n = video_number(name)
        if n is None:
            continue
        verdict = str(row[2] or "").strip().lower()
        violation = bool(verdict) and "без" not in verdict
        labels.append(Label(
            n=n, name=name, violation=violation,
            t_sec=parse_timestamp(row[1]) if violation else None,
            type_id=TYPE_BY_N.get(n) if violation else None,
            video=videos.get(n),
        ))
    return sorted(labels, key=lambda l: l.n)
